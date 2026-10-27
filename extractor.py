"""Safe, bounded text extraction for demo artwork uploads."""
import io
import os
import subprocess
import tempfile
import zipfile
import csv

import fitz
from docx import Document
from PIL import Image, UnidentifiedImageError
from pptx import Presentation

MAX_BYTES = 10_000_000
MAX_PARTS = 20
MAX_EXTRACT = 40_000
ACCEPTED = {".png", ".jpg", ".jpeg", ".webp", ".pptx", ".pdf", ".docx"}
Image.MAX_IMAGE_PIXELS = 25_000_000


def run_ocr(image):
    if image.width * image.height > 25_000_000:
        raise ValueError("Image exceeds pixel limit")
    image = image.convert("RGB")
    with tempfile.NamedTemporaryFile(suffix=".png") as f:
        image.save(f.name)
        result = subprocess.run(
            ["tesseract", f.name, "stdout", "-l", "tha+eng", "--psm", "11"],
            capture_output=True, text=True, timeout=25,
        )
        if result.returncode:
            raise ValueError("OCR failed")
        return result.stdout.strip()


def image_words(image):
    with tempfile.NamedTemporaryFile(suffix=".png") as f:
        image.convert("RGB").save(f.name)
        result = subprocess.run(["tesseract", f.name, "stdout", "-l", "tha+eng", "tsv"],
                                text=True, capture_output=True, timeout=25)
        if result.returncode:
            return []
        boxes = []
        for row in csv.DictReader(io.StringIO(result.stdout), delimiter="\t"):
            token = (row.get("text") or "").strip()
            if not token:
                continue
            try:
                x, y, w, h = (int(row[k]) for k in ("left","top","width","height"))
                boxes.append({"page":1,"text":token,"x":x/image.width,"y":y/image.height,
                              "w":w/image.width,"h":h/image.height})
            except (ValueError,KeyError,ZeroDivisionError):
                continue
        return boxes[:700]


def extract_picture(data):
    image = Image.open(io.BytesIO(data))
    image.load()
    return run_ocr(image)


def safe_zip(data):
    z = zipfile.ZipFile(io.BytesIO(data))
    infos = z.infolist()
    if len(infos) > 3000 or sum(i.file_size for i in infos) > 100_000_000:
        z.close()
        raise ValueError("Document exceeds decompressed size limits")
    return z


def extract_file(filename, data):
    ext = os.path.splitext(filename.lower())[1]
    if ext not in ACCEPTED:
        raise ValueError("Unsupported type. Use PNG, JPG, WEBP, PPTX, PDF, DOCX. Legacy .ppt must be saved as .pptx.")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File must be between 1 byte and 10 MB")
    out = []
    notes = []
    regions = []
    pages = []
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        image = Image.open(io.BytesIO(data))
        image.load()
        out.append(run_ocr(image))
        regions = image_words(image)
        pages = [{"number":1,"kind":"image"}]
    elif ext == ".pptx":
        z = safe_zip(data)
        try:
            prs = Presentation(io.BytesIO(data))
            if len(prs.slides) > MAX_PARTS:
                notes.append("Only first 20 slides evaluated")
            for i, slide in enumerate(prs.slides, 1):
                if i > MAX_PARTS:
                    break
                chunks = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        chunks.append(shape.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            chunks.extend(c.text for c in row.cells)
                    if shape.shape_type == 13:
                        try:
                            chunks.append(extract_picture(shape.image.blob))
                        except Exception:
                            notes.append("Some image text on slide %d could not be read" % i)
                out.append("[Slide %d]\\n%s" % (i, "\\n".join(chunks)))
                pages.append({"number":i,"kind":"slide"})
                for sh in slide.shapes:
                    if sh.has_text_frame and sh.text.strip():
                        regions.append({"page":i,"text":sh.text[:400],
                          "x":float(sh.left/prs.slide_width),"y":float(sh.top/prs.slide_height),
                          "w":float(sh.width/prs.slide_width),"h":float(sh.height/prs.slide_height)})
        finally:
            z.close()
        notes.append("PPTX extraction does not verify layout, visibility, or all grouped/embedded visual content")
    elif ext == ".pdf":
        pdf = fitz.open(stream=data, filetype="pdf")
        try:
            if pdf.is_encrypted:
                raise ValueError("Password-protected PDF unsupported")
            if len(pdf) > MAX_PARTS:
                notes.append("Only first 20 pages evaluated")
            for i in range(min(len(pdf), MAX_PARTS)):
                p = pdf[i]
                text = p.get_text()
                if len(text.strip()) < 30:
                    pix = p.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    image = Image.open(io.BytesIO(pix.tobytes("png")))
                    text += "\n" + run_ocr(image)
                out.append("[Page %d]\\n%s" % (i + 1, text))
                pages.append({"number":i+1,"kind":"pdf"})
                for w in p.get_text("words")[:1000]:
                    x0,y0,x1,y1,t,*rest=w
                    regions.append({"page":i+1,"text":t,"x":x0/p.rect.width,
                       "y":y0/p.rect.height,"w":(x1-x0)/p.rect.width,"h":(y1-y0)/p.rect.height})
        finally:
            pdf.close()
    elif ext == ".docx":
        z = safe_zip(data)
        try:
            doc = Document(io.BytesIO(data))
            out.extend(p.text for p in doc.paragraphs)
            for table in doc.tables[:50]:
                for row in table.rows[:100]:
                    out.append(" | ".join(c.text for c in row.cells))
            for rel in list(doc.part.rels.values())[:200]:
                if "image" in rel.reltype:
                    try:
                        out.append(extract_picture(rel.target_part.blob))
                    except Exception:
                        notes.append("Some embedded images could not be read")
        finally:
            z.close()
        notes.append("DOCX extraction may miss headers, text boxes and complex layouts")
    text = "\n".join(filter(None, out)).strip()
    if len(text) > MAX_EXTRACT:
        text = text[:MAX_EXTRACT]
        notes.append("Text truncated to 40,000 characters")
    if not text:
        notes.append("No readable text found; human review required")
    return {"text": text, "status": "extracted" if text else "not_evaluated", "warnings": notes, "file_type": ext, "units_examined": min(len(out), MAX_PARTS), "regions":regions[:1600], "pages":pages[:MAX_PARTS]}
